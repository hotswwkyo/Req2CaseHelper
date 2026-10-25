"""
Convert MD test cases to Excel format.

Usage:
    python md2excel.py
    python md2excel.py file1.md file2.md
    python md2excel.py -o ./output
    python md2excel.py --force
    python md2excel.py --with-module     # 将MD中"所属模块"转入Excel对应列
    python md2excel.py --title-prefix-module   # 在"用例标题"前加前缀【所属模块】
    python md2excel.py --split 90        # 按条数分页：每文件最多90条，多页加 _001/_002 后缀
    python md2excel.py --split-module    # 按"所属模块"拆分：每个模块一个文件
"""

import argparse
import json
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

TESTCASES_DIR = Path(__file__).parent.parent.parent.parent.parent / "testcases"
RECORD_FILE = Path(__file__).parent.parent.parent.parent.parent / "md2excel_record.json"

EXCEL_COLUMNS = [
    ("所属模块", None, ""),
    ("相关需求", None, ""),
    ("用例标题", "用例标题", None),
    ("前置条件", "前置条件", None),
    ("步骤", "测试步骤", None),
    ("预期", "预期结果", None),
    ("关键词", None, ""),
    ("优先级", None, "2"),
    ("用例类型", None, "功能测试"),
    ("适用阶段", None, ""),
    ("用例状态", None, "正常"),
]

DEFAULT_PAGE_SIZE = 90          # --split 缺省时的每文件条数
GROUP_UNASSIGNED = "未分组"      # 按模块拆分时空模块的归组名


def load_record():
    """加载转换记录；旧格式/结构不符直接视为空记录（旧数据不迁移）。"""
    if not RECORD_FILE.exists():
        return []
    try:
        data = json.loads(RECORD_FILE.read_text(encoding="utf-8"))
    except Exception:
        return []
    if not isinstance(data, list):
        return []
    for item in data:
        if (not isinstance(item, dict) or "md" not in item
                or "excels" not in item or not isinstance(item["excels"], list)):
            return []
    return data


def save_record(record):
    RECORD_FILE.write_text(json.dumps(record, ensure_ascii=False, indent=2), encoding="utf-8")


def update_record(record, md_path, excel_paths):
    """替换式更新：md 的记录被设为本次实际生成的 excel 路径列表。"""
    paths = list(dict.fromkeys(excel_paths))
    for item in record:
        if item["md"] == md_path:
            item["excels"] = paths
            return
    record.append({"md": md_path, "excels": paths})


def is_already_converted(record, md_path, excel_path):
    return any(item["md"] == md_path and excel_path in item["excels"] for item in record)


def parse_md_table(table_text):
    lines = [tt.strip() for tt in table_text.strip().split("\n") if tt.strip()]
    if len(lines) < 2:
        return []
    h = [c.strip() for c in lines[0].split("|") if c.strip()]
    rows = []
    for line in lines[2:]:
        cells = [c.strip() for c in line.split("|") if c.strip()]
        if len(cells) != len(h):
            continue
        rows.append(dict(zip(h, cells)))
    return rows


def find_tables_in_md(content):
    import re as _re
    tables = []
    for m in _re.finditer(r"^(\|.+\|\s*\n(?:\|.*\|.*\n)*)", content, _re.MULTILINE):
        rows = parse_md_table(m.group(1))
        if rows:
            tables.append(rows)
    return tables


def find_relevant_tables(tables):
    required = {"用例标题", "前置条件", "测试步骤", "预期结果"}
    return [r for r in tables if r and required.issubset(set(r[0].keys()))]


def convert_rows(rows, with_module=False, title_prefix_module=False):
    result = []
    for row in rows:
        module_name = str(row.get("所属模块", "") or "").strip()
        er = []
        for en, md_name, fv in EXCEL_COLUMNS:
            # "所属模块"特殊处理：with_module时从MD映射，否则用默认空值
            if en == "所属模块" and with_module:
                er.append(row.get("所属模块", ""))
            elif en == "用例标题" and title_prefix_module and module_name:
                # 用例标题加模块前缀：【所属模块】标题
                title = format_cell_value(en, row.get(md_name, ""))
                er.append(f"【{module_name}】{title}")
            elif fv is not None:
                er.append(fv)
            elif md_name:
                er.append(format_cell_value(en, row.get(md_name, "")))
            else:
                er.append("")
        result.append(er)
    return result


def format_steps(raw_text: str) -> str:
    """
    将原始步骤文本按分号拆分成多行，保留编号。
    示例输入: "1. 点击【新增用户】；2. 填写所有必填字段（如用户名=zhangsan、姓名=张三）；3. 办公电话输入"13800138000"；4. 点击【保存】"
    输出:
        1. 点击【新增用户】
        2. 填写所有必填字段（如用户名=zhangsan、姓名=张三）
        3. 办公电话输入"13800138000"
        4. 点击【保存】
    """
    # 按中文分号拆分，并去掉每个部分首尾空格
    steps = [step.strip() for step in raw_text.split('；') if step.strip()]
    # 如果分号是英文分号，也可以兼容（用户输入可能是英文分号）
    if len(steps) == 1 and ';' in raw_text:
        steps = [step.strip() for step in raw_text.split(';') if step.strip()]
    # 用换行符连接
    return '\n'.join(steps)


def format_expected(raw_text: str) -> str:
    """
    处理预期结果文本，按分号拆分成多行，没有编号则增加编号。
    """
    # 按中文分号拆分，并去掉每个部分首尾空格
    steps = [step.strip() for step in raw_text.split('；') if step.strip()]
    # 如果分号是英文分号，也可以兼容（用户输入可能是英文分号）
    if len(steps) == 1 and ';' in raw_text:
        steps = [step.strip() for step in raw_text.split(';') if step.strip()]

    steps = [
        f"{i+1}. {step}" if not step.startswith(f"{i+1}.") else step for i, step in enumerate(steps)]
    # 用换行符连接
    return '\n'.join(steps)


def format_cell_value(col_name, value):
    if col_name == "步骤":
        return format_steps(value)
    elif col_name == "预期":
        return format_expected(value)
    else:
        return value


def write_excel(filepath, all_rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "测试用例"
    hf = Font(name="微软雅黑 Light", size=11, bold=True)
    df = Font(name="微软雅黑 Light", size=11)
    hfill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    tb = Border(left=Side(style="thin"), right=Side(style="thin"),
                top=Side(style="thin"), bottom=Side(style="thin"))
    for ci, (en, _, _) in enumerate(EXCEL_COLUMNS, 1):
        c = ws.cell(row=1, column=ci, value=en)
        c.font = hf
        c.fill = hfill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = tb
    for ri, rd in enumerate(all_rows, 2):
        for ci, v in enumerate(rd, 1):
            c = ws.cell(row=ri, column=ci, value=v)
            c.font = df
            c.alignment = Alignment(wrap_text=True, vertical="center")
            c.border = tb
    cm = {n: i for i, (n, _, _) in enumerate(EXCEL_COLUMNS)}
    dw = {"所属模块": 14, "相关需求": 14, "用例标题": 24, "前置条件": 30, "步骤": 50,
          "预期": 50, "关键词": 12, "优先级": 10, "用例类型": 12, "适用阶段": 10, "用例状态": 10}
    for en, ci in cm.items():
        ws.column_dimensions[openpyxl.utils.get_column_letter(ci + 1)].width = dw.get(en, 16)
    wb.save(filepath)


def sanitize_filename_part(part: str) -> str:
    """模块名转安全文件名：替换非法字符与空白，去除首尾分隔符。"""
    import re as _re
    s = _re.sub(r'[\\/:*?"<>|]', "_", part)
    s = _re.sub(r"\s+", "_", s)
    s = s.strip("_.")
    return s or GROUP_UNASSIGNED


def chunk_rows(rows, size):
    return [rows[i:i + size] for i in range(0, len(rows), size)]


def group_rows_by_module(rows):
    """按"所属模块"（首列）保持首次出现顺序分组；空模块归入 '未分组'。"""
    groups = {}
    for row in rows:
        m = (row[0] or "").strip() if row else ""
        key = m or GROUP_UNASSIGNED
        groups.setdefault(key, []).append(row)
    return list(groups.items())


def resolve_unique_names(base_names):
    """重名文件名追加 _2/_3... 消歧（保序、保扩展名）。"""
    seen, out = {}, []
    for bn in base_names:
        seen[bn] = seen.get(bn, 0) + 1
        if seen[bn] == 1:
            out.append(bn)
        else:
            p = Path(bn)
            out.append(f"{p.stem}_{seen[bn]}{p.suffix}")
    return out


def plan_outputs(md_path, rows, split, split_module, has_module_col):
    """按拆分模式规划输出，返回 (chunks, file_names)。

    - 不拆分/退化单文件：仅 1 个 chunk，文件名 = 原默认名（无后缀）。
    - 按条数分页：多页时文件名 `stem_001/_002...`。
    - 按模块拆分：多组时文件名 `stem_<模块>`；无「所属模块」列则退化为单文件。
    """
    stem = md_path.stem
    base = md_path.with_suffix(".xlsx").name
    if split_module:
        if not has_module_col:
            return [rows], [base]
        groups = group_rows_by_module(rows)
        if len(groups) == 1:
            return [[r for _, rs in groups for r in rs]], [base]
        names = [f"{stem}_{sanitize_filename_part(m)}.xlsx" for m, _ in groups]
        return [rs for _, rs in groups], resolve_unique_names(names)
    if split is not None:
        chunks = chunk_rows(rows, split)
        if len(chunks) == 1:
            return chunks, [base]
        names = [f"{stem}_{i + 1:03d}.xlsx" for i in range(len(chunks))]
        return chunks, names
    return [rows], [base]


def process_md_file(md_path, output_dir, record, force, with_module=False,
                    split=None, split_module=False, title_prefix_module=False):
    """转换单个 MD；返回本次新生成的 Excel 文件数（0 = 未生成/已转换）。"""
    tables = find_relevant_tables(find_tables_in_md(md_path.read_text(encoding="utf-8")))
    if not tables:
        return 0
    eff_module = with_module or split_module
    all_rows = []
    for t in tables:
        all_rows.extend(convert_rows(t, eff_module, title_prefix_module))
    if not all_rows:
        return 0
    has_module_col = any("所属模块" in t[0] for t in tables)
    if title_prefix_module and not has_module_col:
        print(f"  [提示] 无「所属模块」列，标题前缀未生效: {md_path.name}")
    chunks, names = plan_outputs(md_path, all_rows, split, split_module, has_module_col)
    if split_module and not has_module_col:
        print(f"  [提示] 无「所属模块」列，按模块拆分退化: {md_path.name} → {names[0]}")
    paths = [output_dir / n for n in names]
    made = 0
    for chunk, p in zip(chunks, paths):
        if not force and is_already_converted(record, str(md_path), str(p)):
            print(f"  [跳过]（已转换）: {p.name}")
            continue
        write_excel(p, chunk)
        print(f"  [已生成]: {p.name}")
        made += 1
    update_record(record, str(md_path), [str(p) for p in paths])
    return made


def _is_int(s):
    try:
        int(s)
        return True
    except (TypeError, ValueError):
        return False


def normalize_split_argv(argv):
    """裸 `--split`（未跟数字）时自动补默认值，避免 argparse 误吞后续文件参数。"""
    out = []
    i = 0
    while i < len(argv):
        tok = argv[i]
        if tok == "--split":
            out.append(tok)
            nxt = argv[i + 1] if i + 1 < len(argv) else None
            if nxt is None or nxt.startswith("-") or not _is_int(nxt):
                out.append(str(DEFAULT_PAGE_SIZE))
        else:
            out.append(tok)
        i += 1
    return out


def main():
    p = argparse.ArgumentParser(description="MD测试用例转Excel（支持分页/按模块拆分到多个文件）")
    p.add_argument("files", nargs="*")
    p.add_argument("-o", "--output", default=None)
    p.add_argument("--force", action="store_true",
                   help="强制重新转换（忽略已转换记录；仅作用于本次处理到的文件，其它文件记录保留）")
    p.add_argument("--with-module", action="store_true",
                   help="将MD中的'所属模块'字段转入Excel对应列（默认不转）")
    p.add_argument("--title-prefix-module", action="store_true",
                   help="在'用例标题'前加前缀【所属模块】（需MD含'所属模块'列）")
    p.add_argument("--split", nargs="?", const=DEFAULT_PAGE_SIZE, type=int, default=None,
                   help=f"按条数分页拆分，每个文件最多 N 条用例（缺省 N={DEFAULT_PAGE_SIZE}）")
    p.add_argument("--split-module", action="store_true",
                   help="按'所属模块'拆分，每个模块一个文件（自动开启 --with-module）")
    args = p.parse_args(normalize_split_argv(sys.argv[1:]))
    if args.split is not None and args.split < 1:
        print("[错误] --split 的 N 必须 >= 1")
        return
    if args.split is not None and args.split_module:
        print("[错误] --split 与 --split-module 不能同时使用")
        return
    mfs = [Path(f) for f in args.files if Path(f).exists()
           ] if args.files else sorted(TESTCASES_DIR.glob("*.md"))
    if not mfs:
        print("没有可转换的 .md 文件")
        return
    od = Path(args.output).resolve() if args.output else None
    if od:
        od.mkdir(parents=True, exist_ok=True)
    rec = load_record()  # --force 只影响"是否跳过已转换"，不清空整份记录
    print(f"[文件夹] 找到 {len(mfs)} 个 .md 文件，开始转换...\n")
    made, skipped = 0, 0
    for mf in mfs:
        n = process_md_file(mf, od or mf.parent, rec, args.force,
                            args.with_module, args.split, args.split_module,
                            args.title_prefix_module)
        if n:
            made += n
        else:
            skipped += 1
    save_record(rec)
    print(f"\n[完成] 转换完成！新生成 {made} 个 Excel，跳过 {skipped} 个 MD（无表/空/已转换）")
    if od:
        print(f"[文件夹] 输出目录: {od}")
    if args.with_module:
        print("[提示] 已启用 --with-module，'所属模块'已转入Excel对应列")
    if args.split is not None:
        print(f"[提示] 已按条数分页，每文件最多 {args.split} 条")
    if args.split_module:
        print("[提示] 已按'所属模块'拆分，每个模块一个文件")
    if args.title_prefix_module:
        print("[提示] 已启用 --title-prefix-module，'用例标题'已加前缀【所属模块】")


if __name__ == "__main__":
    main()
